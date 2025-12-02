import sqlite3
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, redirect, render_template, request, url_for
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
DB_PATH = BASE_DIR / "database.db"

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 512 * 1024 * 1024  # 512 MB limit
ALLOWED_EXCEL_SUFFIXES = {".xlsx"}


def init_storage():
    UPLOAD_DIR.mkdir(exist_ok=True)
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS uploads (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                folder_name TEXT NOT NULL,
                relative_path TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                size_bytes INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL
            )
            """
        )
        conn.commit()


def sanitize_relative_path(path: str) -> Path:
    cleaned_parts = []
    for part in Path(path).parts:
        if part in {"..", ".", ""}:
            continue
        cleaned_parts.append(part)
    return Path(*cleaned_parts)


def sanitize_folder_label(label: str) -> str:
    cleaned = "".join(ch for ch in label if ch.isalnum() or ch in {" ", "_", "-"}).strip()
    return cleaned or "Ordner"


def save_file(file_storage, folder_label: str):
    relative = sanitize_relative_path(file_storage.filename)
    if not relative.parts:
        return None

    destination = UPLOAD_DIR / folder_label / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    file_storage.save(destination)

    size = destination.stat().st_size
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            INSERT INTO uploads (folder_name, relative_path, stored_path, size_bytes, uploaded_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                folder_label,
                str(relative),
                str(destination.relative_to(BASE_DIR)),
                size,
                datetime.utcnow().isoformat(timespec="seconds"),
            ),
        )
        conn.commit()
    return relative


def fetch_recent_uploads():
    with sqlite3.connect(DB_PATH) as conn:
        return conn.execute(
            """
            SELECT folder_name, relative_path, stored_path, size_bytes, uploaded_at
            FROM uploads
            ORDER BY uploaded_at DESC
            LIMIT 25
            """
        ).fetchall()


def render_home(*, names=None, success_message=None, error_message=None):
    return render_template(
        "index.html",
        uploads=fetch_recent_uploads(),
        names=names or [],
        success_message=success_message,
        error_message=error_message,
    )


def parse_excel_names(file_storage):
    suffix = Path(file_storage.filename or "").suffix.lower()
    if suffix not in ALLOWED_EXCEL_SUFFIXES:
        raise ValueError("Ungültiger Dateityp. Bitte eine .xlsx-Datei hochladen.")

    file_storage.stream.seek(0)
    try:
        workbook = load_workbook(file_storage, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Die Datei konnte nicht gelesen werden. Bitte das offizielle Template verwenden.") from exc

    sheet = workbook.active
    header_rows = sheet.iter_rows(max_row=1, values_only=False)
    try:
        first_row = next(header_rows)
    except StopIteration as exc:  # empty sheet
        raise ValueError("Leere Datei. Bitte das offizielle Template verwenden.") from exc

    header = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in first_row]
    if len(header) < 2 or header[0] != "vorname" or header[1] != "nachname":
        raise ValueError("Ungültiges Format – bitte das offizielle Template verwenden.")

    names = []
    for first_name, last_name in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
        if first_name is None and last_name is None:
            continue
        first_clean = str(first_name).strip() if first_name is not None else ""
        last_clean = str(last_name).strip() if last_name is not None else ""
        if not first_clean and not last_clean:
            continue
        names.append({"vorname": first_clean, "nachname": last_clean})

    if not names:
        raise ValueError("Keine Namen gefunden. Bitte Zeilen ausfüllen und erneut versuchen.")

    return names


@app.route("/")
def index():
    return render_home()


@app.route("/upload", methods=["POST"])
def upload():
    files = request.files.getlist("files")
    folder_label_input = request.form.get("folderName", "")
    folder_label = sanitize_folder_label(folder_label_input)

    if not files:
        return jsonify({"error": "Keine Dateien erhalten."}), 400

    saved = 0
    for file_storage in files:
        saved_path = save_file(file_storage, folder_label)
        if saved_path:
            saved += 1

    if saved == 0:
        return jsonify({"error": "Es konnte nichts gespeichert werden."}), 400

    return redirect(url_for("index"))


@app.route("/upload-namensliste", methods=["GET", "POST"])
def upload_namensliste():
    if request.method == "GET":
        return render_home()

    file_storage = request.files.get("namensliste")
    if not file_storage or file_storage.filename == "":
        return render_home(error_message="Bitte wähle eine Excel-Datei aus.")

    try:
        names = parse_excel_names(file_storage)
    except ValueError as exc:
        return render_home(error_message=str(exc))
    except Exception:  # noqa: BLE001
        return render_home(
            error_message="Die Namensliste konnte nicht verarbeitet werden. Bitte das offizielle Template verwenden."
        )

    return render_home(
        names=names, success_message="Namensliste wurde erfolgreich hochgeladen und verarbeitet."
    )


init_storage()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
